#!/usr/bin/env python3
"""将中文视频项目的分类 Markdown 源稿导出为易读、可复制的 Word 文档。"""

from __future__ import annotations

import argparse
import io
import os
import re
import sys
sys.dont_write_bytecode = True
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urlparse
from zipfile import BadZipFile, ZipFile

try:
    from PIL import Image, ImageOps, UnidentifiedImageError
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Cm, Inches, Pt, RGBColor
except ImportError as exc:
    raise SystemExit(
        "缺少导出依赖。请按技能说明使用 Codex 工作区依赖中的 "
        "python-docx 与 Pillow 运行本脚本。原始错误："
        f"{exc}"
    ) from exc


FONT_NAME = "Microsoft YaHei"
COLOR_TEXT = "000000"
COLOR_LINK = "0563C1"
COLOR_MUTED = "666666"
COLOR_PROMPT = "F3F5F7"
COLOR_WARNING = "B42318"
MAX_IMAGE_WIDTH_IN = 6.45
MAX_IMAGE_HEIGHT_IN = 8.3

HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")
UNORDERED_RE = re.compile(r"^[ \t]*[-+*][ \t]+(.+)$")
ORDERED_RE = re.compile(r"^[ \t]*(\d+)[.)][ \t]+(.+)$")
IMAGE_RE = re.compile(r"^[ \t]*!\[([^\]]*)\]\((.+)\)[ \t]*$")
FENCE_RE = re.compile(r"^[ \t]*((?:\x60{3,}|~{3,}))[ \t]*([^ \t]*)[ \t]*$")
TOC_RE = re.compile(r"^[ \t]*<!--[ \t]*(?:目录|TOC)[ \t]*-->[ \t]*$", re.I)
PAGE_BREAK_RE = re.compile(r"^[ \t]*<!--[ \t]*(?:分页|PAGE[ -]?BREAK)[ \t]*-->[ \t]*$", re.I)
COMMENT_RE = re.compile(r"^[ \t]*<!--.*-->[ \t]*$")


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    ordered_number: str = ""
    language: str = ""
    destination: str = ""
    title: str = ""
    bookmark: str = ""
    bookmark_id: int = 0


def smart_join(parts: list[str]) -> str:
    if not parts:
        return ""
    result = parts[0].rstrip()
    for raw in parts[1:]:
        current = raw.strip()
        if not current:
            continue
        if result.endswith("  "):
            result = result[:-2] + "\n" + current
            continue
        left = result[-1:] if result else ""
        right = current[:1]
        cjk_left = bool(re.match(r"[\u3000-\u9fff\uff00-\uffef]", left))
        cjk_right = bool(re.match(r"[\u3000-\u9fff\uff00-\uffef]", right))
        separator = "" if cjk_left and cjk_right else " "
        result += separator + current
    return result


def split_destination(raw: str) -> tuple[str, str]:
    value = raw.strip()
    if value.startswith("<"):
        close = value.find(">")
        if close != -1:
            destination = value[1:close]
            rest = value[close + 1 :].strip()
            title = rest[1:-1] if len(rest) >= 2 and rest[0] == rest[-1] and rest[0] in "'\"" else ""
            return destination, title
    match = re.match(r"^(.*?)[ \t]+([\"'])(.*?)\2[ \t]*$", value)
    if match:
        return match.group(1).strip(), match.group(3)
    return value, ""


def strip_frontmatter(lines: list[str]) -> list[str]:
    if not lines or lines[0].strip() != "---":
        return lines
    for index in range(1, min(len(lines), 200)):
        if lines[index].strip() == "---":
            return lines[index + 1 :]
    return lines


def parse_markdown(markdown: str) -> list[Block]:
    lines = strip_frontmatter(markdown.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    blocks: list[Block] = []
    paragraph_lines: list[str] = []
    fence_marker = ""
    fence_language = ""
    fence_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        text = smart_join(paragraph_lines).strip()
        if text:
            blocks.append(Block("paragraph", text=text))
        paragraph_lines = []

    for line in lines:
        if fence_marker:
            closing = re.match(
                rf"^[ \t]*{re.escape(fence_marker[0])}{{{len(fence_marker)},}}[ \t]*$",
                line,
            )
            if closing:
                blocks.append(
                    Block(
                        "code",
                        text="\n".join(fence_lines).rstrip("\n"),
                        language=fence_language,
                    )
                )
                fence_marker = ""
                fence_language = ""
                fence_lines = []
            else:
                fence_lines.append(line)
            continue

        fence = FENCE_RE.match(line)
        if fence:
            flush_paragraph()
            fence_marker = fence.group(1)
            fence_language = fence.group(2)
            fence_lines = []
            continue

        if TOC_RE.match(line):
            flush_paragraph()
            blocks.append(Block("toc"))
            continue
        if PAGE_BREAK_RE.match(line):
            flush_paragraph()
            blocks.append(Block("page_break"))
            continue
        if COMMENT_RE.match(line):
            flush_paragraph()
            continue
        if not line.strip():
            flush_paragraph()
            continue

        heading = HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            blocks.append(
                Block(
                    "heading",
                    text=heading.group(2).strip(),
                    level=len(heading.group(1)),
                )
            )
            continue

        image = IMAGE_RE.match(line)
        if image:
            flush_paragraph()
            destination, title = split_destination(image.group(2))
            blocks.append(
                Block(
                    "image",
                    text=image.group(1).strip(),
                    destination=destination,
                    title=title,
                )
            )
            continue

        ordered = ORDERED_RE.match(line)
        if ordered:
            flush_paragraph()
            blocks.append(
                Block(
                    "list",
                    text=ordered.group(2).strip(),
                    ordered_number=ordered.group(1),
                )
            )
            continue

        unordered = UNORDERED_RE.match(line)
        if unordered:
            flush_paragraph()
            blocks.append(Block("list", text=unordered.group(1).strip()))
            continue

        if line.lstrip().startswith("> "):
            paragraph_lines.append(line.lstrip()[2:])
        else:
            paragraph_lines.append(line)

    if fence_marker:
        blocks.append(
            Block(
                "code",
                text="\n".join(fence_lines).rstrip("\n"),
                language=fence_language,
            )
        )
    flush_paragraph()

    bookmark_counter = 1
    for block in blocks:
        if block.kind == "heading":
            block.bookmark = f"vw_heading_{bookmark_counter}"
            block.bookmark_id = bookmark_counter
            bookmark_counter += 1
    return blocks


def set_xml_font(rpr, font_name: str = FONT_NAME) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)


def configure_style(
    style,
    *,
    size: float,
    bold: bool | None = None,
    before: float = 0,
    after: float = 5,
    keep_with_next: bool = False,
) -> None:
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(COLOR_TEXT)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    set_xml_font(rpr)
    style_ppr = style.element.get_or_add_pPr()
    paragraph_borders = style_ppr.find(qn("w:pBdr"))
    if paragraph_borders is not None:
        style_ppr.remove(paragraph_borders)
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = 1.15
    paragraph_format.keep_with_next = keep_with_next


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    configure_style(document.styles["Normal"], size=11, after=5)
    configure_style(
        document.styles["Title"],
        size=22,
        bold=False,
        before=10,
        after=5,
        keep_with_next=True,
    )
    configure_style(
        document.styles["Heading 1"],
        size=16,
        bold=True,
        before=10,
        after=5,
        keep_with_next=True,
    )
    configure_style(
        document.styles["Heading 2"],
        size=13,
        bold=True,
        before=10,
        after=5,
        keep_with_next=True,
    )
    configure_style(
        document.styles["Heading 3"],
        size=11,
        bold=True,
        before=10,
        after=5,
        keep_with_next=True,
    )

    if "提示词" not in document.styles:
        prompt_style = document.styles.add_style("提示词", WD_STYLE_TYPE.PARAGRAPH)
        prompt_style.base_style = document.styles["Normal"]
    prompt_style = document.styles["提示词"]
    configure_style(prompt_style, size=10.5, after=6)
    prompt_style.paragraph_format.space_before = Pt(3)
    prompt_style.paragraph_format.left_indent = Cm(0.18)
    prompt_style.paragraph_format.right_indent = Cm(0.18)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shading = ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        ppr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def add_text_run(paragraph, text: str, *, bold: bool = False, italic: bool = False, code: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    run.font.name = FONT_NAME
    if code:
        run.font.size = Pt(10.5)
    rpr = run._element.get_or_add_rPr()
    set_xml_font(rpr)
    return run


def add_hyperlink(paragraph, text: str, target: str = "", *, anchor: str = "") -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if anchor:
        hyperlink.set(qn("w:anchor"), anchor)
        hyperlink.set(qn("w:history"), "1")
    else:
        relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    set_xml_font(rpr)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLOR_LINK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    text_element = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERNS = [
    (
        "link",
        re.compile(r"(?<!!)\[([^\]]+)\]\(([^)]+)\)"),
    ),
    (
        "bold",
        re.compile(r"\*\*(.+?)\*\*|__(.+?)__", re.S),
    ),
    (
        "code",
        re.compile(r"\x60([^\x60]+)\x60", re.S),
    ),
    (
        "italic",
        re.compile(r"(?<!\*)\*([^*\n]+)\*(?!\*)|(?<!_)_([^_\n]+)_(?!_)", re.S),
    ),
]


def unescape_markdown(text: str) -> str:
    return re.sub(r"\\([\\*_\[\]()#!])", r"\1", text)


def add_inline_markdown(paragraph, text: str) -> None:
    lines = text.split("\n")
    for line_index, line in enumerate(lines):
        position = 0
        while position < len(line):
            candidates = []
            for kind, pattern in INLINE_PATTERNS:
                match = pattern.search(line, position)
                if match:
                    candidates.append((match.start(), kind, match))
            if not candidates:
                add_text_run(paragraph, unescape_markdown(line[position:]))
                break
            _, kind, match = min(candidates, key=lambda item: (item[0], item[2].end()))
            if match.start() > position:
                add_text_run(paragraph, unescape_markdown(line[position : match.start()]))
            if kind == "link":
                destination, _ = split_destination(match.group(2))
                add_hyperlink(paragraph, match.group(1), destination)
            elif kind == "bold":
                add_text_run(paragraph, match.group(1) or match.group(2), bold=True)
            elif kind == "italic":
                add_text_run(paragraph, match.group(1) or match.group(2), italic=True)
            elif kind == "code":
                add_text_run(paragraph, match.group(1), code=True)
            position = match.end()
        if line_index < len(lines) - 1:
            paragraph.add_run().add_break()


def add_bookmark_start(paragraph, block: Block) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(block.bookmark_id))
    start.set(qn("w:name"), block.bookmark)
    paragraph._p.append(start)


def add_bookmark_end(paragraph, block: Block) -> None:
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(block.bookmark_id))
    paragraph._p.append(end)


def toc_headings(blocks: Iterable[Block]) -> list[Block]:
    headings = [block for block in blocks if block.kind == "heading" and block.level >= 2]
    if not headings:
        headings = [block for block in blocks if block.kind == "heading"]
    top_level = min((block.level for block in headings), default=2)
    return [block for block in headings if block.level == top_level]


def add_static_toc(document: Document, blocks: list[Block]) -> None:
    entries = toc_headings(blocks)
    if not entries:
        return
    if len(entries) <= 6:
        paragraph = document.add_paragraph()
        add_text_run(paragraph, "快速定位：", bold=True)
        for index, entry in enumerate(entries):
            if index:
                add_text_run(paragraph, "　/　")
            add_hyperlink(paragraph, entry.text, anchor=entry.bookmark)
        return

    title = document.add_paragraph()
    run = add_text_run(title, "目录", bold=True)
    run.font.size = Pt(12)
    for entry in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.35)
        add_hyperlink(paragraph, entry.text, anchor=entry.bookmark)


def resolve_local_image(markdown_path: Path, destination: str) -> Path | None:
    parsed = urlparse(destination)
    if parsed.scheme in {"http", "https"}:
        return None
    if parsed.scheme == "file":
        raw_path = unquote(parsed.path)
        if os.name == "nt" and re.match(r"^/[A-Za-z]:", raw_path):
            raw_path = raw_path[1:]
        candidate = Path(raw_path)
    else:
        candidate = Path(unquote(destination))
    if not candidate.is_absolute():
        candidate = markdown_path.parent / candidate
    try:
        return candidate.resolve()
    except OSError:
        return candidate.absolute()


def image_stream_and_size(image_path: Path) -> tuple[io.BytesIO, float, float]:
    with Image.open(image_path) as source:
        normalized = ImageOps.exif_transpose(source)
        normalized.load()
        if normalized.mode not in {"RGB", "RGBA"}:
            normalized = normalized.convert("RGBA" if "A" in normalized.getbands() else "RGB")
        width_px, height_px = normalized.size
        dpi_value = source.info.get("dpi", (144, 144))
        try:
            dpi_x = float(dpi_value[0])
            dpi_y = float(dpi_value[1])
        except (TypeError, ValueError, IndexError):
            dpi_x = dpi_y = 144.0
        if not 36 <= dpi_x <= 1200:
            dpi_x = 144.0
        if not 36 <= dpi_y <= 1200:
            dpi_y = 144.0
        width_in = max(0.5, width_px / dpi_x)
        height_in = max(0.5, height_px / dpi_y)
        scale = min(1.0, MAX_IMAGE_WIDTH_IN / width_in, MAX_IMAGE_HEIGHT_IN / height_in)
        stream = io.BytesIO()
        normalized.save(stream, format="PNG")
        stream.seek(0)
        return stream, width_in * scale, height_in * scale


def add_warning_paragraph(document: Document, message: str) -> None:
    paragraph = document.add_paragraph()
    run = add_text_run(paragraph, message, italic=True)
    run.font.color.rgb = RGBColor.from_string(COLOR_WARNING)


def add_image(
    document: Document,
    markdown_path: Path,
    block: Block,
    warnings: list[str],
) -> None:
    image_path = resolve_local_image(markdown_path, block.destination)
    if image_path is None:
        message = f"未嵌入远程图片：{block.destination}"
        warnings.append(message)
        add_warning_paragraph(document, f"［{message}］")
        return
    if not image_path.is_file():
        message = f"图片不存在：{image_path}"
        warnings.append(message)
        add_warning_paragraph(document, f"［{message}］")
        return
    try:
        stream, width_in, _ = image_stream_and_size(image_path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        shape = paragraph.add_run().add_picture(stream, width=Inches(width_in))
        if block.text:
            shape._inline.docPr.set("descr", block.text)
            shape._inline.docPr.set("title", block.text)
        if block.title:
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = add_text_run(caption, block.title)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(COLOR_MUTED)
    except (OSError, ValueError, UnidentifiedImageError) as exc:
        message = f"图片无法嵌入：{image_path}（{exc}）"
        warnings.append(message)
        add_warning_paragraph(document, f"［{message}］")


def heading_style(level: int, title_already_used: bool) -> str:
    if level == 1 and not title_already_used:
        return "Title"
    if level <= 2:
        return "Heading 1"
    if level == 3:
        return "Heading 2"
    return "Heading 3"


def render_blocks(document: Document, markdown_path: Path, blocks: list[Block]) -> list[str]:
    warnings: list[str] = []
    title_already_used = False

    for block in blocks:
        if block.kind == "heading":
            style = heading_style(block.level, title_already_used)
            if style == "Title":
                title_already_used = True
            paragraph = document.add_paragraph(style=style)
            add_bookmark_start(paragraph, block)
            add_inline_markdown(paragraph, block.text)
            add_bookmark_end(paragraph, block)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, block.text)
        elif block.kind == "list":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.55)
            paragraph.paragraph_format.first_line_indent = Cm(-0.35)
            prefix = f"{block.ordered_number}. " if block.ordered_number else "• "
            add_text_run(paragraph, prefix)
            add_inline_markdown(paragraph, block.text)
        elif block.kind == "code":
            paragraph = document.add_paragraph(style="提示词")
            set_paragraph_shading(paragraph, COLOR_PROMPT)
            if block.text:
                lines = block.text.split("\n")
                for index, line in enumerate(lines):
                    add_text_run(paragraph, line, code=True)
                    if index < len(lines) - 1:
                        paragraph.add_run().add_break()
            else:
                add_text_run(paragraph, " ")
        elif block.kind == "toc":
            add_static_toc(document, blocks)
        elif block.kind == "page_break":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        elif block.kind == "image":
            add_image(document, markdown_path, block, warnings)
    return warnings


def validate_docx(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
            missing = sorted(required.difference(archive.namelist()))
            if missing:
                raise RuntimeError(f"DOCX 缺少必要文件：{', '.join(missing)}")
            bad_member = archive.testzip()
            if bad_member:
                raise RuntimeError(f"DOCX 压缩包损坏：{bad_member}")
        Document(path)
    except (BadZipFile, OSError, ValueError) as exc:
        raise RuntimeError(f"DOCX 校验失败：{exc}") from exc


def build_document(markdown_path: Path) -> tuple[Document, list[str]]:
    markdown = markdown_path.read_text(encoding="utf-8-sig")
    blocks = parse_markdown(markdown)
    document = Document()
    configure_document(document)
    warnings = render_blocks(document, markdown_path, blocks)

    title = next(
        (block.text for block in blocks if block.kind == "heading" and block.level == 1),
        markdown_path.stem,
    )
    document.core_properties.title = re.sub(r"[*_\x60]", "", title)
    document.core_properties.subject = "中文视频工作流文档"
    document.core_properties.author = "Codex"
    document.core_properties.keywords = "视频工作流, TapNow, 项目说明书"
    return document, warnings


def rebase_links(document: Document, markdown_path: Path, output_path: Path) -> None:
    """图片已嵌入；只重算外部本地超链接，兼容源稿与 Word 不同目录。"""
    for relationship in document.part.rels.values():
        if not relationship.is_external or relationship.reltype != RT.HYPERLINK:
            continue
        target = relationship.target_ref
        if target.startswith('#') or urlparse(target).scheme:
            continue
        decoded = unquote(target).replace('\\', '/')
        destination = (markdown_path.parent / decoded).resolve()
        relationship._target = os.path.relpath(destination, output_path.parent).replace('\\', '/')


def save_to_temporary(document: Document, output_path: Path, *, update: bool) -> Path:
    if output_path.exists() and not update:
        raise FileExistsError(
            f"目标文件已存在：{output_path}\n"
            "请先把 Word 中的修改或批注合并进 Markdown，再使用 --更新 重新导出。"
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.",
        suffix=".docx",
        dir=output_path.parent,
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        document.save(temporary_path)
        validate_docx(temporary_path)
        return temporary_path
    except Exception:
        if temporary_path.exists():
            temporary_path.unlink()
        raise


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="将剧本、分镜图提示词、视频提示词与台词或简短说明书导出为同名 Word。",
    )
    parser.add_argument("markdown", type=Path, help="分类文档 Markdown 源稿路径")
    parser.add_argument(
        "--项目目录", "--project", type=Path, dest="project",
        help="显式指定项目根目录；用于新项目或无法从源稿位置可靠识别时",
    )
    parser.add_argument(
        "--草稿", "--draft", action="store_true", dest="draft",
        help="导出预览草稿；缺失内容列为待补，默认写入项目/临时工作/导出预览",
    )
    parser.add_argument(
        "--更新",
        "--update",
        "-u",
        action="store_true",
        dest="update",
        help="确认已合并现有 Word 修改，并允许替换同名 DOCX",
    )
    parser.add_argument(
        "--输出",
        "--output",
        type=Path,
        dest="output",
        help="可选同名输出路径；只允许项目/临时工作或项目/归档/临时工作内",
    )
    parser.add_argument(
        "--修复归档", action="store_true", dest="repair_archive",
        help="仅在已明确授权修复历史归档时使用；不会自动修改链接",
    )
    parser.add_argument(
        "--目录", action="store_true", dest="include_toc",
        help="仅在用户明确需要目录时允许导出源稿中的目录标记；仍须实际验证跳转",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    markdown_path = args.markdown.expanduser().resolve()
    if not markdown_path.is_file():
        print(f"错误：找不到 Markdown 文件：{markdown_path}", file=sys.stderr)
        return 2
    if markdown_path.suffix.lower() not in {".md", ".markdown"}:
        print(f"错误：输入文件不是 Markdown：{markdown_path}", file=sys.stderr)
        return 2

    from 工作路径 import plan_export
    from 内容检查 import check_source_content
    from 检查项目 import check_document_pair, check_source
    temporary_path = None
    try:
        plan = plan_export(
            markdown_path,
            project=args.project,
            output=args.output,
            draft=args.draft,
            repair_archive=args.repair_archive,
        )
        output_path = plan.output
        if not args.include_toc and any(TOC_RE.match(line) for line in markdown_path.read_text(encoding="utf-8-sig").splitlines()):
            raise RuntimeError("默认不生成目录或镜头导航。请删除目录标记；仅在用户明确需要时使用 --目录。")
        # 草稿允许内容未齐，但路径和链接仍必须安全；正式稿必须先通过内容检查。
        source_problems = check_source(
            markdown_path, output_path, root=plan.project,
            check_content=not plan.draft,
        )
        if plan.draft:
            # 草稿不是历史/正式配对，忽略“只能作为草稿导出”这一元数据提示。
            source_problems = [p for p in source_problems if "只能作为草稿导出" not in p]
        if source_problems:
            raise RuntimeError("；".join(source_problems))
        content = check_source_content(markdown_path, formal=not plan.draft)
        if content["问题"]:
            raise RuntimeError("内容检查未通过：" + "；".join(content["问题"]))
        document, warnings = build_document(markdown_path)
        if warnings:
            raise RuntimeError("；".join(warnings))
        rebase_links(document, markdown_path, output_path)
        temporary_path = save_to_temporary(document, output_path, update=args.update)
        pair_problems = check_document_pair(
            markdown_path, temporary_path, logical_output=output_path,
            root=plan.project, check_content=not plan.draft,
        )
        if plan.draft:
            pair_problems = [p for p in pair_problems if "只能作为草稿导出" not in p]
        if pair_problems:
            raise RuntimeError("导出后检查未通过：" + "；".join(pair_problems))
        # 所有检查完成后才原子替换，任何失败都保留旧稿。
        os.replace(temporary_path, output_path)
        temporary_path = None
    except (FileExistsError, OSError, RuntimeError, UnicodeError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()

    print(output_path)
    for item in content.get("待补", []):
        print(f"待补：{item}", file=sys.stderr)
    for item in content.get("待人工", []):
        print(f"待人工：{item}", file=sys.stderr)
    for warning in warnings:
        print(f"警告：{warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
