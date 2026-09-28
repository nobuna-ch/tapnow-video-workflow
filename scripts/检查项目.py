#!/usr/bin/env python3
"""只读检查视频项目；text 模式仅使用标准库，word 模式另核对 DOCX。"""
from __future__ import annotations

import argparse
from hashlib import sha256
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
from urllib.parse import unquote, urlparse
from zipfile import BadZipFile, ZipFile
import xml.etree.ElementTree as ET

sys.dont_write_bytecode = True
from 工作路径 import (CURRENT_KINDS, canonical_kind, infer_project_root, is_current_source,
                      is_history_source, is_relative_to, is_temporary_target)
from 内容检查 import check_current_content, check_source_content

KINDS = CURRENT_KINDS
VERSION = re.compile(r"(?:" + "|".join(map(re.escape, KINDS)) + r")_第(\d{2,})版\.(md|docx)$")
CURRENT = re.compile(r"(?:" + "|".join(map(re.escape, KINDS)) + r")\.(md|docx)$")
LINK = re.compile(r"(!?)\[([^\]]+)\]\(([^)\n]+)\)")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def exporter():
    name = "_video_workflow_exporter_check"
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, Path(__file__).with_name("导出说明书.py"))
        if spec is None or spec.loader is None:
            raise RuntimeError("无法加载导出器")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    return sys.modules[name]


def project_root(path: Path, explicit: Path | None = None) -> Path:
    return infer_project_root(path, explicit)


def history_dirs(root: Path) -> tuple[Path, ...]:
    return root / "归档" / "历史版本", root / "历史版本"


def document_kind(path: Path) -> str:
    return path.stem.rsplit("_第", 1)[0]


def default_output(markdown: Path, explicit_project: Path | None = None) -> Path:
    root = project_root(markdown, explicit_project)
    return root / markdown.with_suffix(".docx").name if is_current_source(markdown, root) else markdown.with_suffix(".docx")


def split_destination(raw: str) -> tuple[str, str]:
    value = raw.strip()
    if value.startswith("<") and ">" in value:
        close = value.find(">")
        return value[1:close], value[close + 1:].strip().strip("'\"")
    match = re.match(r"^(.*?)[ \t]+([\"'])(.*?)\2[ \t]*$", value)
    return (match.group(1).strip(), match.group(3)) if match else (value, "")


def check_metadata(markdown: Path, output: Path | None = None, *, root: Path | None = None) -> list[str]:
    markdown = markdown.resolve()
    try:
        root = (root or project_root(markdown)).resolve()
    except ValueError as exc:
        return [str(exc)]
    output = (output or default_output(markdown, root)).resolve()
    problems: list[str] = []
    if is_current_source(markdown, root):
        if output.name != markdown.with_suffix(".docx").name:
            problems.append("当前 Word 与 Markdown 必须同名")
        if output != root / markdown.with_suffix(".docx").name and not is_temporary_target(output, root):
            problems.append("当前 Word 应在项目根目录或项目临时工作内")
        text = markdown.read_text(encoding="utf-8-sig")
        title = next((line for line in text.splitlines() if line.startswith("# ")), "")
        if markdown.stem not in title:
            problems.append("文档标题未包含对应文档类型")
        if re.search(r"第\d+版", title) or re.search(r"^(版本|版本记录)\s*[:：]", text, re.M):
            problems.append("当前文档不保留版号标题、版本字段或版本记录")
    elif is_history_source(markdown, root):
        if output != markdown.with_suffix(".docx"):
            problems.append("历史 Word 必须与 Markdown 原位同名配对")
    else:
        problems.append("源稿不是当前固定文档或静态历史稿；只能作为草稿导出")
    return problems


def check_target(path: Path, destination: str, label: str = "", *, root: Path | None = None) -> list[str]:
    parsed = urlparse(destination)
    if destination.startswith("#") or parsed.scheme in {"http", "https", "mailto"}:
        return []
    decoded = unquote(destination).replace("\\", "/")
    if decoded.startswith("file:"):
        return [f"本地链接应使用项目内相对路径：{destination}"]
    try:
        root = (root or project_root(path)).resolve()
    except ValueError as exc:
        return [str(exc)]
    target = (path.parent / decoded).resolve()
    problems: list[str] = []
    if not target.is_file():
        problems.append(f"本地链接不存在：{destination}")
    if not is_relative_to(target, root):
        problems.append(f"本地链接位于项目之外：{destination}")
    elif "临时工作" in target.relative_to(root).parts:
        problems.append(f"正式稿不得依赖临时工作文件：{destination}")
    historical = is_history_source(path, root)
    if historical and target.stem in KINDS and target.suffix.lower() in {".md", ".docx"} and not is_history_source(target, root):
        problems.append(f"历史稿不得链接会更新的当前文档：{destination}")
    if historical and not VERSION.fullmatch(target.name) and target.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".mp4", ".mov", ".wav", ".mp3", ".txt", ".md", ".docx"}:
        fixed = (*history_dirs(root), root / "归档" / "原始资料", root / "原始资料", root / "归档" / "方案资料")
        if not any(is_relative_to(target, folder) for folder in fixed):
            problems.append(f"历史素材需使用归档副本或原始资料：{destination}")
    return problems


def check_source(markdown: Path, output: Path | None = None, *, root: Path | None = None,
                 check_content: bool = True) -> list[str]:
    markdown = markdown.resolve()
    try:
        root = (root or project_root(markdown)).resolve()
    except ValueError as exc:
        return [str(exc)]
    output = output or default_output(markdown, root)
    problems = check_metadata(markdown, output, root=root)
    text = markdown.read_text(encoding="utf-8-sig")
    for _, label, raw in LINK.findall(text):
        destination, _ = split_destination(raw)
        problems.extend(check_target(markdown, destination, label, root=root))
    if check_content and is_current_source(markdown, root):
        problems.extend(check_source_content(markdown, formal=True)["问题"])
    return list(dict.fromkeys(problems))


def _paragraphs(document) -> list[str]:
    return ["".join(p._p.xpath(".//w:t/text()")) for p in document.paragraphs if p._p.xpath(".//w:t/text()")]


def _package_facts(source) -> tuple[list[str], list[tuple[str, str]]]:
    with ZipFile(source) as archive:
        image_hashes = sorted(
            sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith("word/media/")
        )
        xml = ET.fromstring(archive.read("word/document.xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        targets = {r.get("Id"): r.get("Target", "") for r in relationships if r.get("TargetMode") == "External"}
        hyperlinks: list[tuple[str, str]] = []
        for hyperlink in xml.iter(W + "hyperlink"):
            destination = targets.get(hyperlink.get(R + "id"))
            if destination:
                label = "".join(n.text or "" for n in hyperlink.iter(W + "t"))
                hyperlinks.append((label, destination))
        return image_hashes, hyperlinks


def check_document_pair(markdown: Path, output: Path | None = None, *, logical_output: Path | None = None,
                        root: Path | None = None, check_content: bool = True) -> list[str]:
    """核对 Word；output 可为尚未替换正式稿的临时文件。"""
    from docx import Document
    markdown = markdown.resolve()
    root = (root or project_root(markdown)).resolve()
    actual_path = (output or default_output(markdown, root)).resolve()
    logical_output = (logical_output or actual_path).resolve()
    problems = check_source(markdown, logical_output, root=root, check_content=check_content)
    if not actual_path.is_file():
        return problems + ["缺少同名 Word"]
    export_module = exporter()
    expected, warnings = export_module.build_document(markdown)
    problems.extend(warnings)
    try:
        actual = Document(actual_path)
        if _paragraphs(expected) != _paragraphs(actual) or len(expected.tables) != len(actual.tables):
            problems.append("Word 与 Markdown 的可见正文不一致")
        export_module.rebase_links(expected, markdown, logical_output)
        expected_stream = io.BytesIO()
        expected.save(expected_stream)
        expected_stream.seek(0)
        expected_images, expected_links = _package_facts(expected_stream)
        actual_images, actual_links = _package_facts(actual_path)
        if expected_images != actual_images:
            problems.append("Word 与 Markdown 的内嵌图片内容不一致")
        if expected_links != actual_links:
            problems.append("Word 与 Markdown 的超链接目标不一致")
        with ZipFile(actual_path) as archive:
            bad = archive.testzip()
            if bad:
                problems.append(f"Word 压缩包损坏：{bad}")
            for label, destination in actual_links:
                problems.extend(check_target(logical_output, destination, label, root=root))
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        problems.append(f"Word 无法读取：{exc}")
    return list(dict.fromkeys(problems))


def _current_sources(root: Path) -> list[Path]:
    folder = root / "归档" / "文档源稿"
    result = list(folder.glob("*.md")) if folder.is_dir() else []
    result.extend(p for p in root.glob("*.md") if CURRENT.fullmatch(p.name))
    return sorted(result)


def check_archived_pair(markdown: Path, *, root: Path, mode: str) -> list[str]:
    """静态历史只核对旧有配对、ZIP 完整性和冻结链接，不重新套当前内容规范。"""
    word = markdown.with_suffix(".docx")
    problems = check_source(markdown, word, root=root, check_content=False)
    if mode != "word":
        return problems
    if not word.is_file():
        return problems + ["缺少同名历史 Word"]
    try:
        with ZipFile(word) as archive:
            bad = archive.testzip()
            if bad:
                problems.append(f"历史 Word 压缩包损坏：{bad}")
            for required in ("[Content_Types].xml", "word/document.xml"):
                if required not in archive.namelist():
                    problems.append(f"历史 Word 缺少必要文件：{required}")
    except (BadZipFile, OSError) as exc:
        problems.append(f"历史 Word 无法读取：{exc}")
    return problems


def check_project(root: Path, *, mode: str = "word") -> dict:
    root = root.resolve()
    problems: list[str] = []
    manual: list[str] = []
    sources = _current_sources(root)
    if not sources:
        problems.append("未找到可检查的当前 Markdown 源稿")
    grouped: dict[str, list[Path]] = {}
    for md in sources:
        grouped.setdefault(canonical_kind(md.stem), []).append(md)
    for name, candidates in grouped.items():
        if len(candidates) > 1:
            problems.append(f"{name}：当前 Markdown 只能保留一份，同类新旧名称不能并存（{'、'.join(p.name for p in candidates)}）")
    for md in sources:
        if not CURRENT.fullmatch(md.name):
            problems.append(f"当前源稿名称不规范：{md.name}")
            continue
        word = root / md.with_suffix(".docx").name
        errors = (check_document_pair(md, word, root=root, check_content=False) if mode == "word"
                  else check_source(md, word, root=root, check_content=False))
        problems.extend(f"{md.name}：{error}" for error in errors)

    words = sorted(p for p in root.glob("*.docx") if CURRENT.fullmatch(p.name))
    if mode == "word" and not words:
        problems.append("根目录缺少当前正式 Word")
    source_names = {md.with_suffix(".docx").name for md in sources if CURRENT.fullmatch(md.name)}
    for word in words:
        if word.name not in source_names:
            problems.append(f"{word.name}：根目录当前 Word 缺少唯一同名 Markdown 源稿")
    allowed = {p.name for p in words} | {p.name for p in sources if p.parent == root}
    for path in root.iterdir():
        if path.is_file() and not path.name.startswith("~$") and path.name not in allowed:
            problems.append(f"根目录有非当前交付文件：{path.name}")

    # 历史仅查旧有配对/链接，不追溯新的逐镜内容规范。
    for folder in history_dirs(root):
        if not folder.is_dir():
            continue
        for md in folder.glob("*.md"):
            if not VERSION.fullmatch(md.name):
                problems.append(f"历史完整稿命名不规范：{md.name}")
            else:
                errors = check_archived_pair(md, root=root, mode=mode)
                problems.extend(f"{md.name}：{error}" for error in errors)

    content = check_current_content(root)
    problems.extend(content["问题"]); manual.extend(content["待人工"])
    skill_root = Path(__file__).resolve().parent.parent
    if any(skill_root.rglob("__pycache__")) or any(skill_root.rglob("*.pyc")):
        problems.append("skill 目录存在运行缓存")
    boundary = ("text 模式只读检查路径、Markdown/TXT、逐镜字段、时间与本地链接；"
                "机器检查不能证明图像一致、口型同步、声音效果、实际节点绑定或 Word 页面效果。")
    if mode == "word":
        boundary = "另核对 Word/Markdown 正文、图片和链接；" + boundary
    return {"通过": not problems, "模式": mode, "当前源稿": [p.name for p in sources],
            "问题": list(dict.fromkeys(problems)), "待人工": list(dict.fromkeys(manual)), "检查边界": boundary}


def document_content(path: Path) -> tuple:
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
        words = tuple(n.text or "" for n in xml.iter(W + "t"))
        images = sorted(sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith("word/media/"))
    return words, images


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project", type=Path)
    parser.add_argument("--模式", "--mode", choices=("text", "word"), default="word", dest="mode")
    args = parser.parse_args(argv)
    if not args.project.is_dir():
        parser.error("项目目录不存在")
    result = check_project(args.project, mode=args.mode)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["通过"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
